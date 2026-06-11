"""Document text extraction for Telegram uploads."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any


@dataclass
class ExtractedDocument:
    filename: str
    mime_type: str | None
    text: str
    status: str = "ok"
    detail: str | None = None


def extract_document_text(filename: str, content: bytes, *, mime_type: str | None = None) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf" or mime_type == "application/pdf":
            return ExtractedDocument(filename, mime_type, _extract_pdf(content))
        if suffix == ".docx" or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return ExtractedDocument(filename, mime_type, _extract_docx(content))
        if suffix == ".pptx" or mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            return ExtractedDocument(filename, mime_type, _extract_pptx(content))
        if suffix == ".xlsx" or mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return ExtractedDocument(filename, mime_type, _extract_xlsx(content))
        if suffix in {
            ".txt",
            ".md",
            ".csv",
            ".json",
            ".yaml",
            ".yml",
            ".toml",
            ".xml",
            ".rtf",
            ".log",
            ".sql",
            ".py",
            ".js",
            ".jsx",
            ".ts",
            ".tsx",
            ".html",
            ".css",
        }:
            return ExtractedDocument(filename, mime_type, content.decode("utf-8", errors="replace"))
    except Exception as exc:
        return ExtractedDocument(filename, mime_type, "", status="failed", detail=str(exc))
    return ExtractedDocument(
        filename,
        mime_type,
        "",
        status="unsupported",
        detail=(
            "Supported text extraction: PDF, DOCX, PPTX, XLSX, TXT, Markdown, CSV, JSON, "
            "YAML, TOML, XML, SQL, logs, and source files."
        ),
    )


def document_memory_body(document: ExtractedDocument, *, max_chars: int = 12000) -> str:
    text = document.text.strip()
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n[Truncated for memory indexing.]"
    return (
        f"Filename: {document.filename}\n"
        f"MIME type: {document.mime_type or 'unknown'}\n"
        f"Extraction status: {document.status}\n"
        f"Detail: {document.detail or 'none'}\n\n"
        "## Extracted Text\n\n"
        f"{text or '[No extractable text found.]'}"
    )


def document_summary(document: ExtractedDocument, *, max_chars: int = 700) -> str:
    text = " ".join(document.text.split())
    if not text:
        return document.detail or "No extractable text found."
    return text[:max_chars]


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    pages = [(page.extract_text() or "").strip() for page in reader.pages[:25]]
    return "\n\n".join(page for page in pages if page)


def _extract_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_pptx(content: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(BytesIO(content))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides[:50], start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        if texts:
            slides.append(f"Slide {index}\n" + "\n".join(text for text in texts if text))
    return "\n\n".join(slides)


def _extract_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets[:20]:
        rows: list[str] = []
        for row in sheet.iter_rows(max_row=200, values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                rows.append(" | ".join(values))
        if rows:
            sections.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
    workbook.close()
    return "\n\n".join(sections)
